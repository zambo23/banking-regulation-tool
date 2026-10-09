"""Excel registers in data/00-source: one workbook per source, one sheet per table, first row is the header."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def read_sheet(path: Path, sheet: str) -> list[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True)
    if sheet not in wb.sheetnames:
        return []
    rows = wb[sheet].iter_rows(values_only=True)
    header = next(rows, None) or []
    return [dict(zip(header, r)) for r in rows if any(v is not None for v in r)]


def write_register(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    """Write {sheet name: (columns, rows)} to path, replacing the workbook."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c) for c in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(columns, start=1):
            width = max((len(str(row.get(col) or "")) for row in rows), default=0)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, len(col)) + 2, 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
